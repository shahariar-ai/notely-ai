import io
import os

from docx import Document
from fpdf import FPDF

FONT_DIR = os.path.join(os.path.dirname(os.path.dirname(os.path.abspath(__file__))), "static", "fonts")


def export_txt(title, transcript, summary) -> bytes:
    lines = [f"Title: {title}", "", "=== SUMMARY ===", summary.get("executive_summary", ""), ""]
    lines.append("=== ACTION ITEMS ===")
    for item in summary.get("action_items", []):
        lines.append(f"- {item.get('task')} (Owner: {item.get('owner')})")
    lines += ["", "=== FULL TRANSCRIPT ===", transcript or ""]
    return "\n".join(lines).encode("utf-8")


def export_docx(title, transcript, summary) -> bytes:
    doc = Document()
    doc.add_heading(title, level=1)
    doc.add_heading("Executive Summary", level=2)
    doc.add_paragraph(summary.get("executive_summary", ""))
    doc.add_heading("Key Highlights", level=2)
    for h in summary.get("key_highlights", []):
        doc.add_paragraph(h, style="List Bullet")
    doc.add_heading("Action Items", level=2)
    for item in summary.get("action_items", []):
        doc.add_paragraph(f"{item.get('task')} — Owner: {item.get('owner')}", style="List Bullet")
    doc.add_heading("Full Transcript", level=2)
    doc.add_paragraph(transcript or "")
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _new_pdf() -> FPDF:
    """PDF with Noto Sans (Latin) + Noto Sans Bengali, with text shaping for Bengali conjuncts."""
    pdf = FPDF()
    pdf.add_font("NotoSans", "", os.path.join(FONT_DIR, "NotoSans-Regular.ttf"))
    pdf.add_font("NotoSans", "B", os.path.join(FONT_DIR, "NotoSans-Bold.ttf"))
    pdf.add_font("NotoBengali", "", os.path.join(FONT_DIR, "NotoSansBengali-Regular.ttf"))
    pdf.add_font("NotoBengali", "B", os.path.join(FONT_DIR, "NotoSansBengali-Bold.ttf"))
    pdf.set_fallback_fonts(["NotoBengali"])
    try:
        pdf.set_text_shaping(True)
    except Exception:  # uharfbuzz missing: still renders, conjuncts may look unshaped
        pass
    pdf.add_page()
    return pdf


def export_pdf(title, transcript, summary) -> bytes:
    pdf = _new_pdf()

    def block(text, style="", size=11, h=6):
        pdf.set_font("NotoSans", style, size)
        pdf.multi_cell(0, h, text or " ", new_x="LMARGIN", new_y="NEXT")

    block(title, "B", 16, 10)
    block("Executive Summary", "B", 12, 8)
    block(summary.get("executive_summary", ""))
    pdf.ln(4)
    block("Action Items", "B", 12, 8)
    for item in summary.get("action_items", []):
        block(f"- {item.get('task')} (Owner: {item.get('owner')})")
    pdf.ln(4)
    block("Full Transcript", "B", 12, 8)
    block(transcript, "", 10)
    return bytes(pdf.output())
